import io

import pytest
from docx import Document

from app.services.resume.extraction import assess_quality, build_profile, extract_contacts
from app.services.resume.parsing import ResumeFileError, extract_text


def test_hh_resume_profile(resume_text):
    profile = build_profile(resume_text)
    names = {s.lower() for s in profile.skill_names()}
    assert {"python", "fastapi", "postgresql", "docker"} <= names
    assert profile.total_experience_months == 16
    assert profile.min_salary == 90_000
    assert profile.level in ("junior", "intern")
    assert any("Junior Python" in p for p in profile.desired_positions)
    assert profile.experience and profile.experience[0].company
    assert profile.education
    # Facts extracted locally are marked with their source; nothing is invented.
    assert all(s.source == "extracted" and not s.confirmed for s in profile.skills)


def test_contacts_extracted_separately(resume_text):
    contacts = extract_contacts(resume_text)
    assert contacts.emails == ["anna.ivanova@example.com"]
    assert contacts.phones


def test_docx_parsing():
    doc = Document()
    doc.add_paragraph("Ключевые навыки")
    doc.add_paragraph("Python, SQL, Git")
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("cv.docx", buf.getvalue())
    assert "Python" in text


def test_extension_must_match_content():
    with pytest.raises(ResumeFileError):
        extract_text("cv.pdf", b"not a pdf at all")
    with pytest.raises(ResumeFileError):
        extract_text("cv.exe", b"MZ")
    with pytest.raises(ResumeFileError):
        extract_text("cv.txt", b"\x00\x01binary")


def test_empty_resume_rejected():
    with pytest.raises(ResumeFileError):
        extract_text("cv.txt", b"   \n  ")


def test_minimal_resume_quality_flags_gaps():
    profile = build_profile("Иван\nХочу работать программистом")
    quality = assess_quality(profile)
    assert quality["score"] < 50
    assert quality["recommendations"]
