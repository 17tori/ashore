"""Safe text extraction from uploaded resumes (PDF, DOCX, TXT)."""

import io
import zipfile

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.errors import AppError

MAX_PDF_PAGES = 30
MAX_TEXT_CHARS = 200_000
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ResumeFileError(AppError):
    code = "invalid_resume_file"


def _ext(filename: str) -> str:
    name = filename.lower()
    return name[name.rfind(".") :] if "." in name else ""


def detect_kind(filename: str, data: bytes) -> str:
    """Trust file content, not the client: magic bytes must agree with the extension."""
    ext = _ext(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeFileError("Поддерживаются только PDF, DOCX и TXT")
    if ext == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ResumeFileError("Файл не похож на PDF")
        return "pdf"
    if ext == ".docx":
        if not data.startswith(b"PK\x03\x04"):
            raise ResumeFileError("Файл не похож на DOCX")
        return "docx"
    if b"\x00" in data[:4096]:
        raise ResumeFileError("Текстовый файл содержит двоичные данные")
    return "txt"


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ResumeFileError("PDF защищён паролем")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ResumeFileError(f"В PDF больше {MAX_PDF_PAGES} страниц")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except PdfReadError as exc:
        raise ResumeFileError("Не удалось прочитать PDF") from exc


def _docx_text(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            # Zip-bomb guard: refuse archives that expand to something unreasonable.
            if sum(i.file_size for i in zf.infolist()) > 50 * 1024 * 1024:
                raise ResumeFileError("DOCX слишком большой после распаковки")
        document = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ResumeFileError("Не удалось прочитать DOCX") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _txt_text(data: bytes) -> str:
    for encoding in ("utf-8", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ResumeFileError("Не удалось определить кодировку текста")


def extract_text(filename: str, data: bytes) -> str:
    kind = detect_kind(filename, data)
    text = {"pdf": _pdf_text, "docx": _docx_text, "txt": _txt_text}[kind](data)
    text = text.replace("\r\n", "\n").replace("\xa0", " ")
    if not text.strip():
        raise ResumeFileError("В файле не найден текст. Если это скан, загрузите текстовую версию")
    return text[:MAX_TEXT_CHARS]
